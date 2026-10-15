"""Build receipt configurations from spreadsheet rows.

A mapping JSON tells receipt-gen which spreadsheet column holds which value.
Each spreadsheet row becomes one receipt. See docs/spreadsheet.md.
"""

from __future__ import annotations

import csv
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from .generator import (
    TAX_LABELS,
    ConfigError,
    load_json,
    receipt_total,
    validate_config,
    validate_receipt,
)


SOURCE_TYPES = ("csv", "xlsx", "google_sheets")
MAPPING_KEYS = {"_comment", "issuer", "source", "include_if", "receipt", "items"}
SOURCE_KEYS = {"type", "path", "url", "worksheet", "header_row", "encoding"}
RECEIPT_KEYS = {"file_id", "name", "issue_date", "amount_note", "remarks"}
ITEM_KEYS = {"name", "quantity", "unit_price", "tax_category", "tax_rate", "tax_amount"}
COLUMN_SPEC_KEYS = {"column", "default", "prefix", "suffix"}
INCLUDE_IF_KEYS = {"column", "in", "not_in"}
DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日")
TRUE_STRINGS = {"TRUE", "YES", "Y", "✓", "○", "〇"}
FALSE_STRINGS = {"FALSE", "NO", "N", "×", "-"}
TAX_CATEGORY_BY_LABEL = {label: key for key, label in TAX_LABELS.items()}
GOOGLE_SCOPES = (
    "https://www.googleapis.com/auth/spreadsheets.readonly",
    "https://www.googleapis.com/auth/drive.readonly",
)


@dataclass
class Table:
    """Spreadsheet values below the header row.

    rows holds (spreadsheet row number, cell values) pairs so that error
    messages can point to the row a person sees in the spreadsheet.
    """

    headers: list[str]
    rows: list[tuple[int, list[Any]]]


@dataclass
class SheetResult:
    config: dict[str, Any]
    base_dir: Path
    skipped: list[str]


def is_mapping(data: Any) -> bool:
    """Return True for a spreadsheet mapping, False for a receipt JSON."""
    return isinstance(data, dict) and "source" in data


def load_mapping(mapping_path: Path) -> dict[str, Any]:
    mapping = load_json(mapping_path)
    validate_mapping(mapping)
    return mapping


def build_config_from_mapping(
    mapping_path: Path,
    source_override: str | dict[str, Any] | None = None,
    gspread_client: Any = None,
    worksheet: str | None = None,
) -> SheetResult:
    """Read the spreadsheet named in a mapping JSON and build a receipt config.

    source_override replaces the mapping's source. A string is treated as a
    Google Sheets URL (when it starts with http) or a local CSV/XLSX path;
    a dict is merged into the mapping's source. worksheet, when given,
    selects the worksheet to read after the override is applied.
    """
    mapping_path = Path(mapping_path).expanduser().resolve()
    mapping = load_mapping(mapping_path)
    base_dir = mapping_path.parent
    source = _apply_source_override(mapping["source"], source_override)
    if worksheet is not None:
        source = {**source, "worksheet": worksheet}
    _validate_source(source)
    table = load_table(source, base_dir, gspread_client=gspread_client)
    config, skipped = build_config(mapping, table)
    return SheetResult(config=config, base_dir=base_dir, skipped=skipped)


# ---------------------------------------------------------------------------
# Reading spreadsheets


def load_table(
    source: dict[str, Any], base_dir: Path, gspread_client: Any = None
) -> Table:
    source_type = source["type"]
    if source_type == "csv":
        values = _read_csv(_source_path(source, base_dir), source.get("encoding"))
    elif source_type == "xlsx":
        values = _read_xlsx(_source_path(source, base_dir), source.get("worksheet"))
    else:
        values = _read_google_sheet(
            source["url"], source.get("worksheet"), gspread_client
        )
    return _make_table(values, source.get("header_row", 1))


def google_client() -> Any:
    """Create a gspread client from Application Default Credentials.

    In Google Colab, call google.colab.auth.authenticate_user() first.
    """
    try:
        import google.auth
        import gspread
    except ImportError as error:
        raise RuntimeError(
            "Reading Google Sheets requires gspread and google-auth. "
            "Use the Colab notebook (docs/colab.md), or export the sheet as "
            "CSV/XLSX and pass it with --source."
        ) from error

    try:
        credentials, _ = google.auth.default(scopes=list(GOOGLE_SCOPES))
    except google.auth.exceptions.DefaultCredentialsError as error:
        raise RuntimeError(
            "No Google credentials found. In Colab, run "
            "google.colab.auth.authenticate_user() first. Locally, export the "
            "sheet as CSV/XLSX and pass it with --source."
        ) from error
    return gspread.authorize(credentials)


def _source_path(source: dict[str, Any], base_dir: Path) -> Path:
    return (base_dir / Path(source["path"]).expanduser()).resolve()


def _read_csv(path: Path, encoding: str | None) -> list[list[Any]]:
    try:
        with path.open(encoding=encoding or "utf-8-sig", newline="") as file:
            return [row for row in csv.reader(file)]
    except FileNotFoundError as error:
        raise ConfigError(f"Spreadsheet file not found: {path}") from error
    except UnicodeDecodeError as error:
        raise ConfigError(
            f"Could not decode {path.name} as {encoding or 'utf-8'}. "
            'For a CSV saved by Excel, set "encoding": "cp932" in source.'
        ) from error


def _read_xlsx(path: Path, worksheet: str | None) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Reading XLSX files requires openpyxl. "
            "Run: python -m pip install -r requirements.txt"
        ) from error

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError as error:
        raise ConfigError(f"Spreadsheet file not found: {path}") from error
    try:
        if worksheet is None:
            sheet = workbook.worksheets[0]
        elif worksheet in workbook.sheetnames:
            sheet = workbook[worksheet]
        else:
            names = ", ".join(workbook.sheetnames)
            raise ConfigError(f'Worksheet "{worksheet}" not found. Available: {names}')
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_google_sheet(
    url: str, worksheet: str | None, client: Any
) -> list[list[Any]]:
    if client is None:
        client = google_client()

    try:
        if url.startswith(("http://", "https://")):
            spreadsheet = client.open_by_url(url)
        else:
            spreadsheet = client.open_by_key(url)
        sheet = spreadsheet.worksheet(worksheet) if worksheet else spreadsheet.sheet1
        return sheet.get_all_values()
    except Exception as error:
        name = type(error).__name__
        if name == "SpreadsheetNotFound":
            raise ConfigError(
                f"Spreadsheet not found or not shared with this account: {url}"
            ) from error
        if name == "WorksheetNotFound":
            raise ConfigError(f'Worksheet "{worksheet}" not found in {url}') from error
        if name == "NoValidUrlKeyFound":
            raise ConfigError(f"Not a Google Sheets URL: {url}") from error
        raise


def _make_table(values: list[list[Any]], header_row: int) -> Table:
    if len(values) < header_row:
        raise ConfigError(f"The spreadsheet has no header row at row {header_row}.")
    headers = [_normalize_header(cell) for cell in values[header_row - 1]]
    rows = []
    for row_number, row in enumerate(values[header_row:], start=header_row + 1):
        row = list(row) + [None] * (len(headers) - len(row))
        rows.append((row_number, row))
    return Table(headers=headers, rows=rows)


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(unicodedata.normalize("NFKC", str(value)).split())


# ---------------------------------------------------------------------------
# Converting rows to receipts


def build_config(
    mapping: dict[str, Any], table: Table
) -> tuple[dict[str, Any], list[str]]:
    """Convert spreadsheet rows to a receipt configuration.

    Returns the configuration and human-readable notes about skipped rows.
    Blank rows are ignored without a note.
    """
    columns = _ColumnLookup(table.headers)
    columns.check_mapping(mapping)

    receipt_spec = mapping["receipt"]
    receipts: list[dict[str, Any]] = []
    skipped: list[str] = []
    for row_number, values in table.rows:
        if all(_is_blank(value) for value in values):
            continue
        row = _Row(row_number, values, columns)
        name = row.text(receipt_spec["name"], "name")
        who = f"row {row_number}" + (f" ({name})" if name else "")

        include_if = mapping.get("include_if")
        if include_if and not row.matches(include_if):
            skipped.append(f'{who}: excluded by include_if ("{include_if["column"]}")')
            continue

        items = [
            item
            for item_spec in mapping["items"]
            if (item := row.item(item_spec)) is not None
        ]
        if not name:
            if items:
                raise ConfigError(f"row {row_number}: name is blank.")
            skipped.append(f"{who}: name is blank and no items are set")
            continue
        if not items:
            skipped.append(f"{who}: no items with quantity > 0")
            continue

        receipt: dict[str, Any] = {
            "file_id": row.text(receipt_spec.get("file_id"), "file_id"),
            "name": name,
            "issue_date": row.date(receipt_spec["issue_date"], "issue_date"),
            "items": items,
            "_source_row": row_number,
        }
        amount_note = row.text(receipt_spec.get("amount_note"), "amount_note")
        if amount_note:
            receipt["amount_note"] = amount_note
        remarks = row.remarks(receipt_spec.get("remarks"))
        if remarks:
            receipt["remarks"] = remarks
        validate_receipt(receipt, who)
        receipts.append(receipt)

    if not receipts:
        detail = f" ({len(skipped)} row(s) skipped)" if skipped else ""
        raise ConfigError(f"No receipts to generate{detail}.")

    config = {"issuer": mapping["issuer"], "receipts": receipts}
    validate_config(config)
    return config, skipped


def preview_rows(config: dict[str, Any]) -> list[dict[str, Any]]:
    """Summarize receipts as table rows for a dry run."""
    rows = []
    for receipt in config["receipts"]:
        items = ", ".join(
            f'{item["name"]}×{item.get("quantity", 1)}' for item in receipt["items"]
        )
        rows.append(
            {
                "row": receipt.get("_source_row"),
                "name": receipt["name"],
                "issue_date": receipt["issue_date"],
                "items": items,
                "total": receipt_total(receipt),
            }
        )
    return rows


class _ColumnLookup:
    def __init__(self, headers: list[str]):
        self.headers = headers
        self.index: dict[str, int] = {}
        self.duplicates: set[str] = set()
        for position, header in enumerate(headers):
            if not header:
                continue
            if header in self.index:
                self.duplicates.add(header)
            self.index[header] = position

    def find(self, column: str) -> int:
        key = _normalize_header(column)
        if key in self.duplicates:
            raise ConfigError(f'Column "{column}" appears more than once in the header.')
        if key not in self.index:
            available = ", ".join(header for header in self.headers if header)
            raise ConfigError(f'Column "{column}" not found. Available: {available}')
        return self.index[key]

    def check_mapping(self, mapping: dict[str, Any]) -> None:
        """Fail early, before any row is read, if a referenced column is missing."""
        for spec in _iter_specs(mapping):
            if isinstance(spec, dict) and "column" in spec:
                self.find(spec["column"])


class _Row:
    def __init__(self, number: int, values: list[Any], columns: _ColumnLookup):
        self.row_number = number
        self.values = values
        self.columns = columns

    def raw(self, spec: Any) -> Any:
        if not isinstance(spec, dict):
            return spec
        value = self.values[self.columns.find(spec["column"])]
        if _is_blank(value):
            return spec.get("default")
        return value

    def where(self, field: str, spec: Any) -> str:
        if isinstance(spec, dict):
            return f'row {self.row_number}, {field} (column "{spec["column"]}")'
        return f"row {self.row_number}, {field}"

    def text(self, spec: Any, field: str) -> str | None:
        value = self.raw(spec)
        if _is_blank(value):
            return None
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        if isinstance(value, datetime):
            value = value.date()
        text = str(value).strip()
        if isinstance(spec, dict):
            prefix = spec.get("prefix", "")
            suffix = spec.get("suffix", "")
            if prefix and not text.startswith(prefix.strip()):
                text = prefix + text
            if suffix and not text.endswith(suffix.strip()):
                text = text + suffix
        return text

    def date(self, spec: Any, field: str) -> str | None:
        value = self.raw(spec)
        if _is_blank(value):
            return None
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        text = unicodedata.normalize("NFKC", str(value)).strip()
        for date_format in DATE_FORMATS:
            try:
                return datetime.strptime(text, date_format).date().isoformat()
            except ValueError:
                continue
        raise ConfigError(
            f"{self.where(field, spec)}: unrecognized date {value!r}. "
            "Use a format such as 2026-07-11 or 2026/07/11."
        )

    def integer(self, spec: Any, field: str, allow_bool: bool = False) -> int | None:
        value = self.raw(spec)
        number = _to_number(value, allow_bool=allow_bool)
        if number is _INVALID or (number is not None and not float(number).is_integer()):
            raise ConfigError(
                f"{self.where(field, spec)}: expected a whole number, got {value!r}."
            )
        return None if number is None else int(number)

    def number(self, spec: Any, field: str) -> int | float | None:
        value = self.raw(spec)
        number = _to_number(value)
        if number is _INVALID:
            raise ConfigError(f"{self.where(field, spec)}: expected a number, got {value!r}.")
        if isinstance(number, float) and number.is_integer():
            number = int(number)
        return number

    def tax_category(self, spec: Any, field: str) -> str | None:
        value = self.text(spec, field)
        if value is None:
            return None
        value = unicodedata.normalize("NFKC", value)
        if value in TAX_LABELS:
            return value
        normalized_labels = {
            unicodedata.normalize("NFKC", label): key
            for label, key in TAX_CATEGORY_BY_LABEL.items()
        }
        if value in normalized_labels:
            return normalized_labels[value]
        choices = ", ".join([*TAX_LABELS, *TAX_CATEGORY_BY_LABEL])
        raise ConfigError(
            f"{self.where(field, spec)}: unknown tax category {value!r}. Use one of: {choices}"
        )

    def item(self, spec: dict[str, Any]) -> dict[str, Any] | None:
        label = f'item "{spec["name"]}"'
        quantity = self.integer(spec["quantity"], f"{label} quantity", allow_bool=True)
        if quantity is None or quantity == 0:
            return None
        if quantity < 0:
            raise ConfigError(
                f'{self.where(f"{label} quantity", spec["quantity"])}: '
                f"quantity must not be negative, got {quantity}."
            )
        unit_price = self.integer(spec["unit_price"], f"{label} unit_price")
        if unit_price is None:
            raise ConfigError(
                f'{self.where(f"{label} unit_price", spec["unit_price"])}: '
                "unit_price is blank."
            )
        item: dict[str, Any] = {
            "name": spec["name"],
            "quantity": quantity,
            "unit_price": unit_price,
        }
        if "tax_category" in spec:
            item["tax_category"] = self.tax_category(
                spec["tax_category"], f"{label} tax_category"
            )
        if "tax_rate" in spec:
            item["tax_rate"] = self.number(spec["tax_rate"], f"{label} tax_rate")
        if "tax_amount" in spec:
            item["tax_amount"] = self.integer(spec["tax_amount"], f"{label} tax_amount")
        return item

    def remarks(self, spec: Any) -> list[str]:
        if spec is None:
            return []
        specs = spec if isinstance(spec, list) else [spec]
        remarks = []
        for remark_spec in specs:
            text = self.text(remark_spec, "remarks")
            if text:
                remarks.append(text)
        return remarks

    def matches(self, include_if: dict[str, Any]) -> bool:
        value = _normalize_match(self.raw({"column": include_if["column"]}))
        if "in" in include_if:
            return value in {_normalize_match(choice) for choice in include_if["in"]}
        return value not in {_normalize_match(choice) for choice in include_if["not_in"]}


_INVALID = object()


def _to_number(value: Any, allow_bool: bool = False) -> Any:
    """Parse cell values such as 31000, 31000.0, "31,000", "¥31,000", "10%".

    Returns None for blanks and _INVALID for values that are not numbers.
    """
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        return int(value) if allow_bool else _INVALID
    if isinstance(value, (int, float)):
        return value
    text = unicodedata.normalize("NFKC", str(value)).strip()
    if allow_bool and text.upper() in TRUE_STRINGS:
        return 1
    if allow_bool and text.upper() in FALSE_STRINGS:
        return 0
    for symbol in (",", "¥", "円", "%", " "):
        text = text.replace(symbol, "")
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        return _INVALID


def _normalize_match(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = unicodedata.normalize("NFKC", str(value)).strip()
    return text.upper() if text.upper() in {"TRUE", "FALSE"} else text


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


# ---------------------------------------------------------------------------
# Validating the mapping itself


def validate_mapping(mapping: Any) -> None:
    if not isinstance(mapping, dict):
        raise ConfigError("The mapping root must be an object.")
    _check_keys(mapping, MAPPING_KEYS, "mapping")
    for key in ("issuer", "source", "receipt", "items"):
        if key not in mapping:
            raise ConfigError(f"mapping.{key} is required.")
    if not isinstance(mapping["issuer"], dict):
        raise ConfigError("mapping.issuer must be an object.")

    _validate_source(mapping["source"])

    receipt = mapping["receipt"]
    if not isinstance(receipt, dict):
        raise ConfigError("mapping.receipt must be an object.")
    _check_keys(receipt, RECEIPT_KEYS, "receipt")
    for key in ("name", "issue_date"):
        if key not in receipt:
            raise ConfigError(f"receipt.{key} is required.")
    for key in ("file_id", "name", "issue_date", "amount_note"):
        if key in receipt:
            _check_spec(receipt[key], f"receipt.{key}")
    remarks = receipt.get("remarks")
    if remarks is not None:
        for index, remark in enumerate(
            remarks if isinstance(remarks, list) else [remarks], start=1
        ):
            _check_spec(remark, f"receipt.remarks[{index}]")

    items = mapping["items"]
    if not isinstance(items, list) or not items:
        raise ConfigError("mapping.items must contain at least one item.")
    for index, item in enumerate(items, start=1):
        label = f"items[{index}]"
        if not isinstance(item, dict):
            raise ConfigError(f"{label} must be an object.")
        _check_keys(item, ITEM_KEYS, label)
        if not isinstance(item.get("name"), str) or not item["name"].strip():
            raise ConfigError(f"{label}.name must be a non-empty string.")
        for key in ("quantity", "unit_price"):
            if key not in item:
                raise ConfigError(f"{label}.{key} is required.")
        for key in ITEM_KEYS - {"name"}:
            if key in item:
                _check_spec(item[key], f"{label}.{key}")

    include_if = mapping.get("include_if")
    if include_if is not None:
        if not isinstance(include_if, dict):
            raise ConfigError("include_if must be an object.")
        _check_keys(include_if, INCLUDE_IF_KEYS, "include_if")
        if not isinstance(include_if.get("column"), str):
            raise ConfigError("include_if.column must be a column name.")
        conditions = [key for key in ("in", "not_in") if key in include_if]
        if len(conditions) != 1 or not isinstance(include_if[conditions[0]], list):
            raise ConfigError('include_if needs exactly one of "in" or "not_in" (a list).')


def _validate_source(source: Any) -> None:
    if not isinstance(source, dict):
        raise ConfigError("mapping.source must be an object.")
    _check_keys(source, SOURCE_KEYS, "source")
    source_type = source.get("type")
    if source_type not in SOURCE_TYPES:
        raise ConfigError(f"source.type must be one of: {', '.join(SOURCE_TYPES)}")
    if source_type in ("csv", "xlsx") and not _is_nonempty_str(source.get("path")):
        raise ConfigError(f'source.path is required for type "{source_type}".')
    if source_type == "google_sheets" and not _is_nonempty_str(source.get("url")):
        raise ConfigError('source.url is required for type "google_sheets".')
    header_row = source.get("header_row", 1)
    if isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 1:
        raise ConfigError("source.header_row must be a positive integer.")
    worksheet = source.get("worksheet")
    if worksheet is not None and not _is_nonempty_str(worksheet):
        raise ConfigError("source.worksheet must be a sheet name or null.")


def _apply_source_override(
    source: dict[str, Any], override: str | dict[str, Any] | None
) -> dict[str, Any]:
    if override is None:
        return source
    if isinstance(override, dict):
        merged = {**source, **override}
        if merged.get("type") != source.get("type") and "worksheet" not in override:
            merged.pop("worksheet", None)
        return merged

    keep = {key: source[key] for key in ("header_row",) if key in source}
    if override.startswith(("http://", "https://")):
        worksheet = source.get("worksheet") if source["type"] == "google_sheets" else None
        return {**keep, "type": "google_sheets", "url": override, "worksheet": worksheet}

    path = Path(override).expanduser().resolve()
    suffix = path.suffix.lower()
    if suffix == ".csv":
        encoding = source.get("encoding") if source["type"] == "csv" else None
        return {**keep, "type": "csv", "path": str(path), "encoding": encoding}
    if suffix in (".xlsx", ".xlsm"):
        worksheet = source.get("worksheet") if source["type"] == "xlsx" else None
        return {**keep, "type": "xlsx", "path": str(path), "worksheet": worksheet}
    raise ConfigError(
        f"Cannot tell the source type of {override}. Use a .csv or .xlsx file "
        "or a Google Sheets URL."
    )


def _check_spec(spec: Any, label: str) -> None:
    if isinstance(spec, dict):
        _check_keys(spec, COLUMN_SPEC_KEYS, label)
        if not _is_nonempty_str(spec.get("column")):
            raise ConfigError(f'{label} needs "column" with a column name.')
    elif isinstance(spec, list):
        raise ConfigError(f"{label} must be a value or an object with \"column\".")


def _check_keys(data: dict[str, Any], allowed: set[str], label: str) -> None:
    unknown = sorted(set(data) - allowed)
    if unknown:
        raise ConfigError(
            f"Unknown key(s) in {label}: {', '.join(unknown)}. "
            f"Allowed: {', '.join(sorted(allowed - {'_comment'}))}"
        )


def _iter_specs(mapping: dict[str, Any]):
    receipt = mapping["receipt"]
    for key in ("file_id", "name", "issue_date", "amount_note"):
        if key in receipt:
            yield receipt[key]
    remarks = receipt.get("remarks")
    if remarks is not None:
        yield from remarks if isinstance(remarks, list) else [remarks]
    for item in mapping["items"]:
        for key in ITEM_KEYS - {"name"}:
            if key in item:
                yield item[key]
    if mapping.get("include_if"):
        yield {"column": mapping["include_if"]["column"]}


def _is_nonempty_str(value: Any) -> bool:
    return isinstance(value, str) and bool(value.strip())
